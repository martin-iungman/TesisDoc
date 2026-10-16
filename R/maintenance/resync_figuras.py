"""Pasos mecánicos de la rutina resync-figuras-mapping (ver docs/rutina_resync_figuras.md).

Un solo punto de entrada para que la tarea programada no tenga que armar
comandos distintos en cada corrida. La parte de juicio (qué slide es qué
slug) sigue siendo de quien corre la rutina; este script solo prepara la
información y lleva el estado del export a Drive.

Uso (desde la raíz del repo):
  python R/maintenance/resync_figuras.py save-deck --title "Fig R2" --src <archivo>
  python R/maintenance/resync_figuras.py changed
  python R/maintenance/resync_figuras.py report [--deck "Fig R2" ...]
  python R/maintenance/resync_figuras.py export-check
  python R/maintenance/resync_figuras.py export-record --file-id <id>

Requiere python-pptx y Pillow.
"""
import argparse
import base64
import binascii
import csv
import datetime
import glob
import hashlib
import io
import json
import os
import re
import subprocess
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DOCS = os.path.join(ROOT, "docs")
FIGURES = os.path.join(ROOT, "figures")
MAPPING = os.path.join(DOCS, "mapping_figuras.csv")
STATE = os.path.join(os.path.expanduser("~"), ".claude", "scheduled-tasks",
                     "resync-figuras-mapping", "export_state.json")
IMG_EXT = (".jpg", ".jpeg", ".png")
# diferencia media por pixel (0-255) en una miniatura 16x16 en gris: por
# debajo de esto se considera la misma imagen recomprimida por Google Slides
NEAR_MAX = 3.0


def deck_path(title):
    return os.path.join(DOCS, title + ".pptx")


def git_changed_decks():
    out = subprocess.run(["git", "status", "--porcelain", "--", "docs/Fig *.pptx"],
                         cwd=ROOT, capture_output=True, text=True, check=True).stdout
    decks = []
    for line in out.splitlines():
        path = line[3:].strip().strip('"')
        decks.append(os.path.splitext(os.path.basename(path))[0])
    return decks


# --- save-deck ---------------------------------------------------------------

def _find_zip_bytes(obj):
    """Busca en un JSON (o string) un campo base64 que decodifique a un zip (pptx)."""
    if isinstance(obj, str):
        s = re.sub(r"\s+", "", obj)
        if len(s) > 100:
            try:
                data = base64.b64decode(s, validate=False)
                if data[:2] == b"PK":
                    return data
            except (binascii.Error, ValueError):
                pass
        return None
    items = obj.values() if isinstance(obj, dict) else obj if isinstance(obj, list) else []
    for v in items:
        found = _find_zip_bytes(v)
        if found:
            return found
    return None


def cmd_save_deck(args):
    raw = open(args.src, "rb").read()
    if raw[:2] == b"PK":
        data = raw
    else:
        text = raw.decode("utf-8", errors="ignore")
        try:
            data = _find_zip_bytes(json.loads(text))
        except json.JSONDecodeError:
            data = _find_zip_bytes(text)
    if not data:
        sys.exit(f"ERROR: no encontré un pptx (base64 de un zip) en {args.src}")
    dst = deck_path(args.title)
    rel = os.path.relpath(dst, ROOT).replace("\\", "/")
    head = subprocess.run(["git", "show", f"HEAD:{rel}"], cwd=ROOT, capture_output=True).stdout
    if head and head != data and _same_content(head, data):
        # Google Slides re-codifica las imágenes en cada export: si texto e
        # imágenes son los mismos, se deja la versión commiteada
        with open(dst, "wb") as fh:
            fh.write(head)
        print(f"{rel}: sin cambios de contenido (solo ruido del export); se deja la versión de HEAD")
        return
    with open(dst, "wb") as fh:
        fh.write(data)
    estado = "nuevo" if not head else ("igual a HEAD" if head == data else "CAMBIÓ respecto de HEAD")
    print(f"{rel}: guardado ({len(data)} bytes, {estado})")


def _fingerprint(data):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    slides = []
    for slide in Presentation(io.BytesIO(data)).slides:
        texts, thumbs = [], []
        for s in _walk(slide.shapes):
            if s.has_text_frame and s.text_frame.text.strip():
                texts.append(s.text_frame.text.strip())
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                thumbs.append(_thumb(s.image.blob))
        slides.append((texts, thumbs))
    return slides


def _same_content(a, b):
    fa, fb = _fingerprint(a), _fingerprint(b)
    if len(fa) != len(fb):
        return False
    for (ta, ia), (tb, ib) in zip(fa, fb):
        if ta != tb or len(ia) != len(ib):
            return False
        for x, y in zip(ia, ib):
            if (x is None) != (y is None):
                return False
            if x is not None and sum(abs(p - q) for p, q in zip(x, y)) / len(x) > 1.0:
                return False
    return True


# --- changed -----------------------------------------------------------------

def cmd_changed(args):
    decks = git_changed_decks()
    if not decks:
        print("sin cambios en los decks de Drive")
    for d in decks:
        print(d)


# --- report ------------------------------------------------------------------

def _walk(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(s.shapes)
        else:
            yield s


def _thumb(data):
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(data)).convert("L").resize((16, 16))
        return list(im.getdata())
    except Exception:
        return None


def _figure_index():
    idx = []
    for p in glob.glob(os.path.join(FIGURES, "**", "*"), recursive=True):
        if os.path.isfile(p) and p.lower().endswith(IMG_EXT):
            d = open(p, "rb").read()
            idx.append((os.path.relpath(p, FIGURES).replace("\\", "/"),
                        hashlib.sha1(d).hexdigest(), _thumb(d)))
    return idx


def _slug_by_number():
    with open(MAPPING, encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    out = {}
    for r in rows:
        m = re.match(r"^Fig\.\s*(M\d+|R\d+\.\d+)$", r["numero_actual"].strip())
        if m:
            out.setdefault(m.group(1), []).append(r["slug"])
    return out


def _deck_prefix(title):
    if title == "Fig MyM":
        return "M"
    m = re.match(r"^Fig R(\d+)$", title)
    return f"R{m.group(1)}." if m else None


def cmd_report(args):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    decks = args.deck or git_changed_decks()
    if not decks:
        print("sin cambios en los decks de Drive")
        return
    idx = _figure_index()
    by_num = _slug_by_number()
    for title in decks:
        path = deck_path(title)
        prefix = _deck_prefix(title)
        print("=" * 70)
        print(f"DECK {title}" + ("" if prefix else "  ** nombre sin esquema de numeración conocido **"))
        for i, slide in enumerate(Presentation(path).slides, 1):
            num = f"{prefix}{i}" if prefix else f"slide {i}"
            slugs = by_num.get(num, []) if prefix else []
            print("-" * 70)
            print(f"SLIDE {i} -> Fig. {num}   (en el CSV hoy: {', '.join(slugs) or 'ningún slug'})")
            texts, pics = [], []
            for s in _walk(slide.shapes):
                if s.has_text_frame and s.text_frame.text.strip():
                    texts.append(s.text_frame.text.strip().replace("\n", " | "))
                if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    pics.append(s.image.blob)
            for t in texts:
                print(f"  texto: {t}")
            if not texts:
                print("  texto: (ninguno)")
            for j, blob in enumerate(pics, 1):
                h, t = hashlib.sha1(blob).hexdigest(), _thumb(blob)
                exact = [p for p, fh_, _ in idx if fh_ == h]
                if exact:
                    print(f"  imagen {j}: IDENTICA a {', '.join(exact)}")
                    continue
                if t is None:
                    print(f"  imagen {j}: no decodificable")
                    continue
                cand = sorted((sum(abs(a - b) for a, b in zip(t, ft)) / len(t), p)
                              for p, _, ft in idx if ft)
                best = [f"{p} ({d:.1f})" for d, p in cand[:2]]
                if not cand or cand[0][0] > NEAR_MAX:
                    tag = "SIN COINCIDENCIA; más parecidas:"
                elif len(cand) > 1 and cand[1][0] - cand[0][0] < 0.5:
                    tag = "AMBIGUA entre"
                else:
                    tag = "COINCIDE con"
                print(f"  imagen {j}: {tag} {'; '.join(best)}")
            if not pics:
                print("  imágenes: (ninguna; formas nativas o placeholder)")


# --- export-check / export-record -------------------------------------------

def _csv_sha():
    return hashlib.sha256(open(MAPPING, "rb").read()).hexdigest()


def _read_state():
    try:
        return json.load(open(STATE, encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return {}


def cmd_export_check(args):
    st, sha = _read_state(), _csv_sha()
    if st.get("sha256") == sha:
        print("UNCHANGED")
    else:
        print("CHANGED")
    print(f"sha256={sha}")
    print(f"previous_file_id={st.get('fileId', '')}")


def cmd_export_record(args):
    os.makedirs(os.path.dirname(STATE), exist_ok=True)
    st = {"sha256": _csv_sha(), "fileId": args.file_id,
          "exportedAt": datetime.date.today().isoformat()}
    with open(STATE, "w", encoding="utf-8") as fh:
        json.dump(st, fh, indent=2)
    print(f"registrado export {st['fileId']} ({st['exportedAt']})")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("save-deck", help="decodifica un pptx bajado de Drive a docs/<title>.pptx")
    p.add_argument("--title", required=True)
    p.add_argument("--src", required=True, help="archivo con el resultado de download_file_content (JSON o base64) o el pptx")
    p.set_defaults(func=cmd_save_deck)
    sub.add_parser("changed", help="decks que difieren del último commit").set_defaults(func=cmd_changed)
    p = sub.add_parser("report", help="texto e imágenes por slide, con coincidencias en figures/")
    p.add_argument("--deck", action="append", help="título del deck (por defecto, los que cambiaron)")
    p.set_defaults(func=cmd_report)
    sub.add_parser("export-check", help="¿cambió el CSV desde el último export a Drive?").set_defaults(func=cmd_export_check)
    p = sub.add_parser("export-record", help="registra el export recién hecho")
    p.add_argument("--file-id", required=True)
    p.set_defaults(func=cmd_export_record)
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    args.func(args)


if __name__ == "__main__":
    main()
